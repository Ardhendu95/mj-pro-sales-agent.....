import io
import math
import re
import zipfile
from collections import Counter
from pathlib import Path
from xml.etree import ElementTree
from sqlalchemy import select
from .database import Document

STOP = set('a an the is are was were i we you my our your it this that to of for in on and or do does can how what me tell about please have with'.split())

def terms(text):
    return Counter(w for w in re.findall(r'[a-z0-9]+', text.lower()) if w not in STOP and len(w) > 1)

def chunks(text, size=1000):
    # Prefer paragraph and sentence boundaries. Never send whole files to the model.
    parts = re.split(r'\n\s*\n|(?<=[.!?])\s+', text)
    output, current = [], ''
    for part in parts:
        for offset in range(0, len(part), size):
            fragment = part[offset:offset + size]
            if len(current) + len(fragment) > size and current:
                output.append(current.strip())
                current = ''
            current += fragment + ' '
    if current.strip():
        output.append(current.strip())
    return output

async def retrieve(db, query):
    docs = (await db.scalars(select(Document).where(Document.approved.is_(True), Document.active.is_(True), Document.confidential.is_(False)))).all()
    q = terms(query)
    scored = []
    for doc in docs:
        for chunk in chunks(doc.text):
            words = terms(chunk + ' ' + doc.title)
            overlap = sum(min(count, words[word]) for word, count in q.items())
            if overlap:
                score = overlap / math.sqrt(max(1, sum(words.values())))
                scored.append((score, doc.created_at, {'id': doc.id, 'title': doc.title, 'source': doc.source, 'version': doc.version, 'excerpt': chunk}))
    return [item[2] for item in sorted(scored, key=lambda x: (x[0], x[1]), reverse=True)[:3]]

def parse_file(filename, content):
    ext = Path(filename).suffix.lower()
    if ext in {'.docx', '.xlsx', '.pptx'}:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(e.file_size for e in archive.infolist()) > 25_000_000:
                raise ValueError('Expanded document is too large. Upload a smaller file.')
    if ext in {'.txt', '.md'}:
        text = content.decode('utf-8-sig')
    elif ext in {'.html', '.htm'}:
        from bs4 import BeautifulSoup
        soup = BeautifulSoup(content, 'html.parser')
        for tag in soup(['script', 'style', 'nav', 'footer']):
            tag.decompose()
        text = soup.get_text(' ', strip=True)
    elif ext == '.pdf':
        from pypdf import PdfReader
        text = '\n\n'.join(p.extract_text() or '' for p in PdfReader(io.BytesIO(content)).pages)
    elif ext == '.docx':
        from docx import Document as Word
        doc = Word(io.BytesIO(content))
        text = '\n'.join([p.text for p in doc.paragraphs] + [' | '.join(c.text for c in row.cells) for table in doc.tables for row in table.rows])
    elif ext == '.xlsx':
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        text = '\n'.join(' | '.join(str(c) for c in row if c is not None) for sheet in book for row in sheet.values)
        book.close()
    elif ext == '.pptx':
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            paths = sorted((p for p in archive.namelist() if re.match(r'ppt/slides/slide\d+\.xml$', p)), key=lambda p: int(re.search(r'(\d+)\.xml', p)[1]))
            text = '\n\n'.join(' '.join(n.text or '' for n in ElementTree.fromstring(archive.read(p)).iter() if n.tag.endswith('}t')) for p in paths)
    else:
        raise ValueError('Use TXT, MD, HTML, PDF, DOCX, XLSX or PPTX.')
    if not text.strip():
        raise ValueError('No readable text found. Scanned PDFs need text recognition first.')
    if len(text) > 300_000:
        raise ValueError('Document contains too much text. Split it into smaller files.')
    return text.strip()
